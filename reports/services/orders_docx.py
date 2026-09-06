import io
from datetime import datetime
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH


def _create_base_order_docx(employee, order_title, body_paragraphs, order_num="1-К", city="Київ"):
    """Базовий шаблон кадрового наказу ФОП згідно зі стандартами діловодства"""
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.78)
        section.bottom_margin = Inches(0.78)
        section.left_margin = Inches(0.78)
        section.right_margin = Inches(0.78)

    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(12)

    fop = getattr(employee, 'fop', None)
    fop_name = fop.full_name if fop else "ФОП Петренко П.П."

    # Шапка
    p_fop = doc.add_paragraph()
    p_fop.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_fop.add_run(f"ФІЗИЧНА ОСОБА-ПІДПРИЄМЕЦЬ\n{fop_name.upper()}\n").bold = True

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(10)
    p_title.paragraph_format.space_after = Pt(10)
    p_title.add_run("НАКАЗ").bold = True

    # Мета-дані
    today = datetime.now()
    p_meta = doc.add_paragraph()
    p_meta.add_run(f"«{today.day}» {today.strftime('%m.%Y')} р.\t\t\tм. {city}\t\t\t№ {order_num}").bold = True

    # Заголовок наказу
    p_subj = doc.add_paragraph()
    p_subj.paragraph_format.space_before = Pt(8)
    p_subj.paragraph_format.space_after = Pt(12)
    p_subj.add_run(f"Про {order_title}").bold = True

    # Текст наказу
    doc.add_paragraph().add_run("НАКАЗУЮ:").bold = True

    for text in body_paragraphs:
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.15
        p.add_run(text)

    # Підписи
    doc.add_paragraph("\n")
    p_sign = doc.add_paragraph()
    p_sign.add_run(f"ФОП\t\t___________________\t\t{fop_name}")

    p_ack = doc.add_paragraph()
    p_ack.paragraph_format.space_before = Pt(16)
    p_ack.add_run(f"З наказом ознайомлений(-а):\n\n___________________\t{employee.full_name}\t«___» ___________ 20__ р.")

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def generate_leave_order_docx(employee, leave_type="щорічну основну відпустку", days=14, start_date=None, order_num="2-К"):
    """Наказ про надання відпустки"""
    date_str = start_date or datetime.now().strftime("%d.%m.%Y")
    body = [
        f"1. НАДАТИ {employee.full_name}, {getattr(employee, 'position', 'працівнику')}, {leave_type} тривалістю {days} календарних днів з {date_str} року.",
        "2. Контроль за виконанням даного наказу залишаю за собою.",
        f"Підстава: заява {employee.full_name} від {date_str} року."
    ]
    return _create_base_order_docx(employee, "надання відпустки", body, order_num=order_num)


def generate_business_trip_order_docx(employee, destination="м. Львів", days=3, start_date=None, order_num="3-К"):
    """Наказ про службове відрядження"""
    date_str = start_date or datetime.now().strftime("%d.%m.%Y")
    body = [
        f"1. НАПРАВИТИ у службове відрядження {employee.full_name}, {getattr(employee, 'position', 'працівника')}, до {destination} терміном на {days} дн. з {date_str} року.",
        "2. Мета відрядження: проведення робочих зустрічей та узгодження виробничих питань.",
        "3. Контроль за виконанням даного наказу залишаю за собою."
    ]
    return _create_base_order_docx(employee, "службове відрядження", body, order_num=order_num)


def generate_salary_change_order_docx(employee, new_salary=25000.0, new_position=None, order_num="4-К"):
    """Наказ про переведення / зміну окладу"""
    pos_text = f" на посаду {new_position}" if new_position else ""
    body = [
        f"1. ЗМІНИТИ {employee.full_name}{pos_text} посадовий оклад та встановити його у розмірі {new_salary:.2f} грн на місяць з «01» числа наступного місяця.",
        "2. Бухгалтерії здійснювати нарахування заробітної плати згідно з новим посадовим окладом.",
        "3. Контроль за виконанням даного наказу залишаю за собою."
    ]
    return _create_base_order_docx(employee, "зміну посадового окладу", body, order_num=order_num)


def generate_financial_aid_order_docx(employee, amount=5000.0, order_num="5-К"):
    """Наказ про надання матеріальної допомоги"""
    body = [
        f"1. НАДАТИ {employee.full_name}, {getattr(employee, 'position', 'працівнику')}, матеріальну допомогу за сімейними обставинами у розмірі {amount:.2f} грн.",
        "2. Виплату матеріальної допомоги здійснити разом із виплатою заробітної плати за поточний місяць.",
        f"Підстава: заява {employee.full_name}."
    ]
    return _create_base_order_docx(employee, "надання матеріальної допомоги", body, order_num=order_num)