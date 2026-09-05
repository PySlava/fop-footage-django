import io
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT


def generate_employment_notice_docx(employee):
    """Генерує Повідомлення про прийняття працівника на роботу у форматі .docx"""
    doc = Document()

    # Поля сторінки (2 см)
    for section in doc.sections:
        section.top_margin = Inches(0.78)
        section.bottom_margin = Inches(0.78)
        section.left_margin = Inches(0.78)
        section.right_margin = Inches(0.78)

    # Дані ФОП
    fop = employee.fop
    fop_name = fop.full_name if fop else "ФОП"
    fop_tax = getattr(fop, 'tax_id', '') if fop else ''

    # Шапка
    p_head = doc.add_paragraph()
    p_head.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_head.add_run(f"До ДПС України\nвід {fop_name}\nІПН/ЄДРПОУ: {fop_tax}\n").bold = True

    # Заголовок
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(12)
    p_title.paragraph_format.space_after = Pt(12)
    run_title = p_title.add_run("ПОВІДОМЛЕННЯ\nпро прийняття працівника на роботу")
    run_title.bold = True
    run_title.font.size = Pt(14)

    # Таблиця з даними працівника
    table = doc.add_table(rows=2, cols=6)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'

    # Заголовки таблиці
    headers = [
        "Категорія особи",
        "Податковий номер (ІПН)",
        "Прізвище, ім'я, по батькові",
        "Номер наказу",
        "Дата видання наказу",
        "Дата початку роботи"
    ]
    hdr_cells = table.rows[0].cells
    for i, header_text in enumerate(headers):
        hdr_cells[i].text = header_text
        hdr_cells[i].paragraphs[0].runs[0].font.bold = True
        hdr_cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Заповнення даними
    hire_date_str = employee.hire_date.strftime("%d.%m.%Y") if hasattr(employee.hire_date, 'strftime') else str(employee.hire_date)
    row_cells = table.rows[1].cells
    row_cells[0].text = "1"  # 1 - Наймані працівники з трудовою книжкою
    row_cells[1].text = str(getattr(employee, 'tax_id', getattr(employee, 'inn', '')))
    row_cells[2].text = str(employee.full_name)
    row_cells[3].text = str(getattr(employee, 'order_num', '1-К'))
    row_cells[4].text = hire_date_str
    row_cells[5].text = hire_date_str

    # Підпис
    doc.add_paragraph("\n\n")
    p_sign = doc.add_paragraph()
    p_sign.add_run(f"Роботодавець: ___________________  / {fop_name} /")

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer