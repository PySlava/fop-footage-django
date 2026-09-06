import io
from datetime import datetime
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT


def generate_labor_contract_docx(employee, contract_num="1", city="Київ"):
    """Генерує трудовий договір між ФОП та працівником у форматі .docx"""
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.78)
        section.bottom_margin = Inches(0.78)
        section.left_margin = Inches(0.78)
        section.right_margin = Inches(0.78)

    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(11)

    fop = getattr(employee, 'fop', None)
    fop_name = fop.full_name if fop else "ФОП Петренко П.П."
    fop_tin = getattr(fop, 'tin', '1234567890') if fop else '1234567890'
    fop_address = getattr(fop, 'address', 'м. Київ, вул. Хрещатик, 1') if fop else 'м. Київ'

    emp_name = employee.full_name
    emp_tax_id = getattr(employee, 'tax_id', '—')
    emp_passport = getattr(employee, 'passport_info', '—')
    hire_date_str = employee.hire_date.strftime("%d.%m.%Y") if hasattr(employee.hire_date, 'strftime') else str(employee.hire_date)
    position = getattr(employee, 'position', 'Фахівець')
    salary = getattr(employee, 'salary', 20000.0)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_run = p_title.add_run(f"ТРУДОВИЙ ДОГОВІР № {contract_num}\n")
    p_run.bold = True
    p_run.font.size = Pt(13)

    p_meta = doc.add_paragraph()
    p_meta.add_run(f"м. {city}").bold = True
    p_meta.add_run(f"\t\t\t\t\t\t\t\t«{datetime.now().day}» {datetime.now().strftime('%m.%Y')} р.")

    p_preamble = doc.add_paragraph()
    p_preamble.paragraph_format.space_before = Pt(8)
    p_preamble.paragraph_format.line_spacing = 1.15
    p_preamble.add_run(
        f"Фізична особа-підприємець {fop_name} (РНОКПП: {fop_tin}), надалі — «Роботодавець», з однієї сторони, "
        f"та громадянин(ка) України {emp_name} (РНОКПП: {emp_tax_id}, паспорт: {emp_passport}), "
        f"надалі — «Працівник», з іншої сторони, уклали цей Договір про наступне:"
    )

    # Розділ 1: Предмет договору
    doc.add_paragraph().add_run("1. ПРЕДМЕТ ДОГОВОРУ").bold = True
    p1 = doc.add_paragraph()
    p1.add_run(f"1.1. Працівник приймається на роботу до Роботодавця на посаду: {position}.\n")
    p1.add_run("1.2. Робота за цим Договором є основним місцем роботи Працівника.\n")
    p1.add_run(f"1.3. Дата початку виконання обов'язків: {hire_date_str}.")

    # Розділ 2: Оплата праці
    doc.add_paragraph().add_run("2. ОПЛАТА ПРАЦІ ТА РЕЖИМ РОБОТИ").bold = True
    p2 = doc.add_paragraph()
    p2.add_run(f"2.1. За виконання обов'язків Працівнику встановлюється посадовий оклад у розмірі {salary:.2f} грн/місяць.\n")
    p2.add_run("2.2. Виплата заробітної плати здійснюється два рази на місяць (аванс та основна виплата).\n")
    p2.add_run("2.3. Працівнику встановлюється 5-денний робочий тиждень з 8-годинним робочим днем.")

    # Розділ 3: Реквізити та підписи сторін
    doc.add_paragraph().add_run("\n3. РЕКВІЗИТИ ТА ПІДПИСИ СТОРІН").bold = True

    table = doc.add_table(rows=2, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'

    hdr_cells = table.rows[0].cells
    hdr_cells[0].paragraphs[0].add_run("РОБОТОДАВЕЦЬ").bold = True
    hdr_cells[1].paragraphs[0].add_run("ПРАЦІВНИК").bold = True

    row_cells = table.rows[1].cells
    row_cells[0].text = f"ФОП {fop_name}\nАдреса: {fop_address}\nРНОКПП: {fop_tin}\n\nПідпис: _____________"
    row_cells[1].text = f"{emp_name}\nПаспорт: {emp_passport}\nРНОКПП: {emp_tax_id}\n\nПідпис: _____________"

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer