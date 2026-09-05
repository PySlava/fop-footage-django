import io
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH


def generate_employee_order_docx(employee):
    """Генерує кадровий наказ про прийняття на роботу у форматі .docx"""
    doc = Document()

    # Поля сторінки (2 см з усіх боків)
    for section in doc.sections:
        section.top_margin = Inches(0.78)
        section.bottom_margin = Inches(0.78)
        section.left_margin = Inches(0.78)
        section.right_margin = Inches(0.78)

    # Дані ФОП
    fop_name = employee.fop.full_name if employee.fop else "ФОП"
    fop_tax = getattr(employee.fop, 'tax_id', '') if employee.fop else ''

    # Шапка документа
    p_fop = doc.add_paragraph()
    p_fop.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_fop.add_run(f"ФІЗИЧНА ОСОБА - ПІДПРИЄМЕЦЬ\n{fop_name.upper()}\n").bold = True
    if fop_tax:
        p_fop.add_run(f"ІПН: {fop_tax}\n")

    # Заголовок
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("\nНАКАЗ\n")
    run_title.bold = True
    run_title.font.size = Pt(16)

    # Дата та номер
    hire_date_str = employee.hire_date.strftime("%d.%m.%Y") if hasattr(employee.hire_date, 'strftime') else str(employee.hire_date)
    p_num = doc.add_paragraph()
    p_num.add_run(f"від {hire_date_str} р.                                                                 № 1-К")

    p_subj = doc.add_paragraph()
    p_subj.add_run("\nПро прийняття на роботу").bold = True

    # Текст наказу
    p_text = doc.add_paragraph()
    p_text.paragraph_format.first_line_indent = Inches(0.5)
    p_text.add_run(
        f"ПРИЙНЯТИ {employee.full_name.upper()} на посаду {employee.position} "
        f"з {hire_date_str} року з посадовим окладом {employee.salary} грн згідно зі штатним розписом."
    )

    # Підписи
    doc.add_paragraph("\n\n")
    p_sign = doc.add_paragraph()
    p_sign.add_run(f"ФОП: ___________________  / {fop_name} /\n\n")
    p_sign.add_run(f"З наказом ознайомлений(-а): ___________________  / {employee.full_name} /")

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer