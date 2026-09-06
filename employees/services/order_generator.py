import io
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

MONTHS_UA = {
    1: "січня", 2: "лютого", 3: "березня", 4: "квітня",
    5: "травня", 6: "червня", 7: "липня", 8: "серпня",
    9: "вересня", 10: "жовтня", 11: "листопада", 12: "грудня"
}


def format_date_ukr(d, full_year=False):
    """
    Перетворює дату у словесно-цифровий формат.
    Приклад: 10 вересня 2026 р. (або 'року', якщо full_year=True)
    """
    if not d:
        return ""
    year_suffix = "року" if full_year else "р."
    return f"{d.day} {MONTHS_UA[d.month]} {d.year} {year_suffix}"


def remove_table_borders(table):
    """Видаляє всі кордони з таблиці docx"""
    tblPr = table._tbl.tblPr
    tblBorders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="none"/>\n'
        f'  <w:left w:val="none"/>\n'
        f'  <w:bottom w:val="none"/>\n'
        f'  <w:right w:val="none"/>\n'
        f'  <w:insideH w:val="none"/>\n'
        f'  <w:insideV w:val="none"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(tblBorders)


def set_cell_margins(cell, top=0, bottom=0, left=0, right=0):
    """Встановлює внутрішні відступи комірки"""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>\n'
        f'  <w:top w:w="{top}" w:type="dxa"/>\n'
        f'  <w:bottom w:w="{bottom}" w:type="dxa"/>\n'
        f'  <w:left w:w="{left}" w:type="dxa"/>\n'
        f'  <w:right w:w="{right}" w:type="dxa"/>\n'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def format_full_name_genitive(full_name):
    """Переводить ПІБ у родовий відмінок (наприклад, Іванова Івана Івановича)"""
    parts = full_name.strip().split()
    if len(parts) < 3:
        return full_name

    last, first, middle = parts[0], parts[1], parts[2]

    # По батькові
    if middle.endswith("вич"):
        middle_gen = middle + "а"
        is_male = True
    elif middle.endswith("вна"):
        middle_gen = middle[:-1] + "и"
        is_male = False
    else:
        middle_gen = middle
        is_male = True

    # Ім'я
    if is_male:
        if first.endswith("о"):
            first_gen = first[:-1] + "а"
        elif first.endswith("й"):
            first_gen = first[:-1] + "я"
        elif not first.endswith(('а', 'я', 'е', 'є', 'и', 'і', 'о', 'у', 'ю')):
            first_gen = first + "а"
        else:
            first_gen = first
    else:
        if first.endswith("а"):
            first_gen = first[:-1] + "и"
        elif first.endswith("я"):
            first_gen = first[:-1] + "ї"
        else:
            first_gen = first

    # Прізвище
    if is_male:
        if last.endswith(('ов', 'єв', 'ев', 'ин', 'ін')):
            last_gen = last + "а"
        elif last.endswith('ий'):
            last_gen = last[:-2] + "ого"
        else:
            last_gen = last
    else:
        if last.endswith('а'):
            last_gen = last[:-1] + "и"
        elif last.endswith('я'):
            last_gen = last[:-1] + "ї"
        else:
            last_gen = last

    return f"{last_gen} {first_gen} {middle_gen}"


def get_initials_genitive(full_name):
    """Повертає прізвище у родовому відмінку з ініціалами (наприклад, Іванова І. І.)"""
    parts = full_name.strip().split()
    if len(parts) >= 3:
        genitive_full = format_full_name_genitive(full_name).split()
        last_gen = genitive_full[0]
        i1 = parts[1][0].upper() + '.'
        i2 = parts[2][0].upper() + '.'
        return f"{last_gen} {i1} {i2}"
    return full_name


def _get_fop_name(employee):
    if hasattr(employee, 'fop') and employee.fop:
        return getattr(employee.fop, 'full_name', str(employee.fop))
    return "Петренко Петро Петрович"


def generate_hiring_order_docx(employee, order_num, order_date, start_date, settlement="м. Київ"):
    """Генерує наказ про призначення / прийняття на роботу (.docx)"""
    doc = Document()

    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(1.5)

    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)

    fop_name = _get_fop_name(employee)
    employee_name = employee.full_name
    position = getattr(employee, 'position', '...') or '...'

    # Шапка
    p_fop = doc.add_paragraph()
    p_fop.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_fop.paragraph_format.space_after = Pt(18)
    run_fop = p_fop.add_run(f"ФОП {fop_name.upper()}")
    run_fop.bold = True
    run_fop.font.size = Pt(12)

    # Назва документа
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(12)
    p_title.paragraph_format.space_after = Pt(18)
    run_title = p_title.add_run("НАКАЗ")
    run_title.bold = True
    run_title.font.size = Pt(14)

    # Мета-рядок
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    remove_table_borders(table)

    col_widths = [Cm(5.0), Cm(6.5), Cm(5.0)]
    row = table.rows[0]
    for idx, width in enumerate(col_widths):
        row.cells[idx].width = width

    # Дата (зліва)
    cell_left = row.cells[0]
    set_cell_margins(cell_left)
    p_left = cell_left.paragraphs[0]
    p_left.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_left.add_run(f"від {format_date_ukr(order_date)}")

    # Населений пункт (посередині)
    cell_mid = row.cells[1]
    set_cell_margins(cell_mid)
    p_mid = cell_mid.paragraphs[0]
    p_mid.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_mid.add_run(settlement if settlement else "....")

    # Номер наказу (справа)
    cell_right = row.cells[2]
    set_cell_margins(cell_right)
    p_right = cell_right.paragraphs[0]
    p_right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_right.add_run(f"№ {order_num}")

    # Тема
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(18)
    p_sub.paragraph_format.space_after = Pt(18)
    p_sub.paragraph_format.line_spacing = 1.15
    p_sub.add_run("Про призначення\n")
    p_sub.add_run(get_initials_genitive(employee_name))

    # НАКАЗУЮ:
    p_cmd = doc.add_paragraph()
    p_cmd.paragraph_format.space_after = Pt(12)
    run_cmd = p_cmd.add_run("НАКАЗУЮ:")
    run_cmd.bold = True

    # Текст
    p_body = doc.add_paragraph()
    p_body.paragraph_format.first_line_indent = Cm(1.25)
    p_body.paragraph_format.space_after = Pt(36)
    p_body.paragraph_format.line_spacing = 1.15
    p_body.alignment = WD_ALIGN_PARAGRAPH.LEFT

    genitive_name = format_full_name_genitive(employee_name)
    pos_str = f"на посаду {position}" if position != "..." else "на посаду ..."
    p_body.add_run(f"ПРИЗНАЧИТИ з {format_date_ukr(start_date, full_year=True)} {genitive_name} {pos_str}")

    # Підписи
    p_sig1 = doc.add_paragraph()
    p_sig1.paragraph_format.space_after = Pt(12)
    p_sig1.add_run(f"ФОП: _______________ / ФОП {fop_name} /")

    p_sig2 = doc.add_paragraph()
    p_sig2.add_run(f"З наказом ознайомлений(а): _______________ / {employee_name} /")

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def generate_dismissal_order_docx(employee, order_num, order_date, dismissal_date, settlement="м. Київ"):
    """Генерує наказ про звільнення (.docx)"""
    doc = Document()

    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(1.5)

    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)

    fop_name = _get_fop_name(employee)
    employee_name = employee.full_name
    position = getattr(employee, 'position', '...') or '...'

    # Шапка
    p_fop = doc.add_paragraph()
    p_fop.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_fop.paragraph_format.space_after = Pt(18)
    run_fop = p_fop.add_run(f"ФОП {fop_name.upper()}")
    run_fop.bold = True
    run_fop.font.size = Pt(12)

    # Назва
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(12)
    p_title.paragraph_format.space_after = Pt(18)
    run_title = p_title.add_run("НАКАЗ")
    run_title.bold = True
    run_title.font.size = Pt(14)

    # Мета-рядок
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    remove_table_borders(table)

    col_widths = [Cm(5.0), Cm(6.5), Cm(5.0)]
    row = table.rows[0]
    for idx, width in enumerate(col_widths):
        row.cells[idx].width = width

    # Дата
    cell_left = row.cells[0]
    set_cell_margins(cell_left)
    p_left = cell_left.paragraphs[0]
    p_left.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_left.add_run(f"від {format_date_ukr(order_date)}")

    # Населений пункт
    cell_mid = row.cells[1]
    set_cell_margins(cell_mid)
    p_mid = cell_mid.paragraphs[0]
    p_mid.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_mid.add_run(settlement if settlement else "....")

    # Номер наказу
    cell_right = row.cells[2]
    set_cell_margins(cell_right)
    p_right = cell_right.paragraphs[0]
    p_right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_right.add_run(f"№ {order_num}")

    # Тема
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(18)
    p_sub.paragraph_format.space_after = Pt(18)
    p_sub.paragraph_format.line_spacing = 1.15
    p_sub.add_run("Про звільнення\n")
    p_sub.add_run(get_initials_genitive(employee_name))

    # НАКАЗУЮ:
    p_cmd = doc.add_paragraph()
    p_cmd.paragraph_format.space_after = Pt(12)
    run_cmd = p_cmd.add_run("НАКАЗУЮ:")
    run_cmd.bold = True

    # Текст
    p_body = doc.add_paragraph()
    p_body.paragraph_format.first_line_indent = Cm(1.25)
    p_body.paragraph_format.space_after = Pt(36)
    p_body.paragraph_format.line_spacing = 1.15
    p_body.alignment = WD_ALIGN_PARAGRAPH.LEFT

    genitive_name = format_full_name_genitive(employee_name)
    pos_str = f"з посади {position}" if position != "..." else "з посади ..."
    p_body.add_run(f"ЗВІЛЬНИТИ {format_date_ukr(dismissal_date, full_year=True)} {genitive_name}, {pos_str} за згодою сторін відповідно до пункту 1 статті 36 КЗпП України.")

    # Підписи
    p_sig1 = doc.add_paragraph()
    p_sig1.paragraph_format.space_after = Pt(12)
    p_sig1.add_run(f"ФОП: _______________ / ФОП {fop_name} /")

    p_sig2 = doc.add_paragraph()
    p_sig2.add_run(f"З наказом ознайомлений(а): _______________ / {employee_name} /")

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer