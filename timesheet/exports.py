import io
import calendar
from django.http import HttpResponse
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

DAY_TYPE_MAP = {
    'WORK': 'Р',
    'WEEKEND': 'В',
    'VACATION': 'ВІ',
    'SICK': 'ТН',
}


def export_timesheet_to_excel(timesheet) -> HttpResponse:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Табель {timesheet.month:02d}-{timesheet.year}"

    # Стилі
    font_bold = Font(name='Arial', size=10, bold=True)
    font_regular = Font(name='Arial', size=10)
    font_title = Font(name='Arial', size=12, bold=True)

    fill_header = PatternFill(start_color="E0E0E0", end_color="E0E0E0", fill_type="solid")
    fill_weekend = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color='A0A0A0'),
        right=Side(style='thin', color='A0A0A0'),
        top=Side(style='thin', color='A0A0A0'),
        bottom=Side(style='thin', color='A0A0A0')
    )

    align_center = Alignment(horizontal='center', vertical='center')
    align_left = Alignment(horizontal='left', vertical='center')

    ws['A1'] = "ТАБЕЛЬ ОБЛІКУ РОБОЧОГО ЧАСУ ТА РОЗРАХУНКОВА ВІДОМІСТЬ"
    ws['A1'].font = font_title
    ws['A3'] = "Працівник:"
    ws['B3'] = str(timesheet.employee)
    ws['A3'].font = font_bold
    ws['B3'].font = font_regular
    ws['A4'] = "Період:"
    ws['B4'] = f"{timesheet.month:02d}/{timesheet.year}"
    ws['A4'].font = font_bold
    ws['B4'].font = font_regular

    start_row = 6

    ws.cell(row=start_row, column=1, value="День").font = font_bold
    ws.cell(row=start_row + 1, column=1, value="Тип").font = font_bold
    ws.cell(row=start_row + 2, column=1, value="Години").font = font_bold

    for r in range(start_row, start_row + 3):
        cell = ws.cell(row=r, column=1)
        cell.fill = fill_header
        cell.border = thin_border
        cell.alignment = align_left

    days_dict = {day.date.day: day for day in timesheet.days.all()}
    _, days_in_month = calendar.monthrange(timesheet.year, timesheet.month)

    for day_num in range(1, days_in_month + 1):
        col = day_num + 1
        day_obj = days_dict.get(day_num)

        day_code = DAY_TYPE_MAP.get(day_obj.day_type, '') if day_obj else ''
        hours = float(day_obj.hours) if day_obj and day_obj.hours > 0 else ''

        c_day = ws.cell(row=start_row, column=col, value=day_num)
        c_day.font = font_bold
        c_day.fill = fill_header

        c_type = ws.cell(row=start_row + 1, column=col, value=day_code)
        c_type.font = font_regular

        c_hours = ws.cell(row=start_row + 2, column=col, value=hours)
        c_hours.font = font_regular

        for cell in (c_day, c_type, c_hours):
            cell.border = thin_border
            cell.alignment = align_center
            if day_code == 'В':
                cell.fill = fill_weekend

    tot_col = days_in_month + 2
    norm_col = days_in_month + 3

    ws.cell(row=start_row, column=tot_col, value="Всього").font = font_bold
    ws.cell(row=start_row + 2, column=tot_col, value=float(timesheet.total_hours)).font = font_bold

    ws.cell(row=start_row, column=norm_col, value="Норма").font = font_bold
    ws.cell(row=start_row + 2, column=norm_col, value=float(timesheet.norm_hours)).font = font_bold

    for col in (tot_col, norm_col):
        for r in range(start_row, start_row + 3):
            cell = ws.cell(row=r, column=col)
            cell.border = thin_border
            cell.alignment = align_center
            if r == start_row:
                cell.fill = fill_header

    salary = getattr(timesheet, 'salary_calculation', None)
    if salary:
        s_row = start_row + 5
        ws.cell(row=s_row, column=1, value="РОЗРАХУНОК ЗАРПЛАТИ ТА ПОДАТКІВ").font = font_title

        salary_rows = [
            ("Оклад (Base)", float(salary.base_salary)),
            ("Нараховано (Gross)", float(salary.gross_salary)),
            ("ПДФО (18%)", float(salary.pdfo)),
            ("Військовий збір (5%)", float(salary.vz)),
            ("ЄСВ (22%, сплачує ФОП)", float(salary.esv)),
            ("До виплати (Net)", float(salary.net_salary)),
        ]

        for idx, (label, val) in enumerate(salary_rows):
            curr_r = s_row + 1 + idx
            lbl_cell = ws.cell(row=curr_r, column=1, value=label)
            val_cell = ws.cell(row=curr_r, column=2, value=val)

            lbl_cell.font = font_bold if "До виплати" in label else font_regular
            val_cell.font = font_bold if "До виплати" in label else font_regular
            val_cell.number_format = '#,##0.00'

            lbl_cell.border = thin_border
            val_cell.border = thin_border

    ws.column_dimensions['A'].width = 24
    ws.column_dimensions['B'].width = 15
    for col in range(2, norm_col + 1):
        col_letter = get_column_letter(col)
        if col > 2:
            ws.column_dimensions[col_letter].width = 4.5

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"timesheet_{timesheet.employee.id}_{timesheet.month:02d}_{timesheet.year}.xlsx"
    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response